#!/usr/bin/env python3
"""엔카 승용 엑셀 2개를 공통 정규화 JSON 하나로 만든다."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import uuid
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

EXPECTED = {
    "manufacturers": 63,
    "modelGroups": 663,
    "generations": 1_256,
    "fuelDrives": 2_158,
    "grades": 5_976,
    "subgrades": 3_297,
}
SNAPSHOT_DATE = "2026-10-04"
NAMESPACE = uuid.UUID("43b3228e-b813-54c3-9c77-5e32ffb1ff42")
REVIEW_MARKERS = ("🔴", "🟡")


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m")
    return str(value).strip()


def source_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m")
    return str(value)


def integer(value: Any) -> int | None:
    if value is None or text(value) == "":
        return None
    return int(float(value))


def nullable_text(value: Any) -> str | None:
    return text(value) or None


def money(value: Any) -> int | None:
    return integer(value)


def ym(value: Any) -> str | None:
    raw = text(value)
    if not raw:
        return None
    if len(raw) == 7 and raw[4] in (".", "-"):
        return f"{raw[:4]}-{raw[5:]}"
    return raw


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def token(*parts: str) -> str:
    return hashlib.sha256("^@".join(parts).encode("utf-8")).hexdigest()[:20]


def stable_uuid(key: str) -> str:
    return str(uuid.uuid5(NAMESPACE, key))


def review(reason: str, extra: str = "") -> tuple[str, str | None]:
    reasons = [item for item in (reason, extra) if item]
    joined = " / ".join(reasons)
    required = any(marker in joined for marker in REVIEW_MARKERS)
    return ("REVIEW_REQUIRED" if required else "CONFIRMED", joined or None)


def rows_by_header(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[sheet_name]
    iterator = sheet.iter_rows(values_only=True)
    headers = [text(value) for value in next(iterator)]
    rows = [dict(zip(headers, values)) for values in iterator if any(value is not None for value in values)]
    workbook.close()
    return rows


def unique_listing_count(rows: Iterable[dict[str, Any]], column: str, description: str) -> int | None:
    values = {integer(row[column]) for row in rows if integer(row[column]) is not None}
    if len(values) > 1:
        raise ValueError(f"{description} 매물 수가 경로 안에서 다릅니다: {sorted(values)}")
    return next(iter(values), None)


def group(rows: list[dict[str, Any]], key_columns: tuple[str, ...]) -> dict[tuple[str, ...], list[dict[str, Any]]]:
    grouped: dict[tuple[str, ...], list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        key = tuple(text(row[column]) for column in key_columns)
        grouped[key].append(row)
    return grouped


def exact_map(rows: list[dict[str, Any]], columns: tuple[str, ...], description: str) -> dict[tuple[str, ...], dict[str, Any]]:
    result: dict[tuple[str, ...], dict[str, Any]] = {}
    for row in rows:
        key = tuple(text(row[column]) for column in columns)
        if key in result:
            raise ValueError(f"{description} 중복 경로: {' › '.join(key)}")
        result[key] = row
    return result


def orders_with_blanks_last(
    rows: list[dict[str, Any]],
    parent_columns: tuple[str, ...],
    name_column: str,
) -> dict[tuple[str, ...], int]:
    """원본 정렬값은 보존하고 빈 값만 같은 부모의 마지막 순번으로 배정한다."""
    parent_max: dict[tuple[str, ...], int] = defaultdict(int)
    for row in rows:
        name = text(row[name_column])
        if not name:
            continue
        parent = tuple(text(row[column]) for column in parent_columns)
        order = integer(row["엔카 정렬순서"])
        if order is not None:
            parent_max[parent] = max(parent_max[parent], order)

    result: dict[tuple[str, ...], int] = {}
    blank_counts: dict[tuple[str, ...], int] = defaultdict(int)
    for row in rows:
        name = text(row[name_column])
        if not name:
            continue
        parent = tuple(text(row[column]) for column in parent_columns)
        order = integer(row["엔카 정렬순서"])
        if order is None:
            blank_counts[parent] += 1
            order = parent_max[parent] + blank_counts[parent]
        result[(*parent, name)] = order
    return result


def load_body_types(path: Path) -> dict[str, dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    if len(rows) != 40:
        raise ValueError(f"현대 바디타입은 40행이어야 합니다: {len(rows)}")
    return {text(row["model_group"]): {"bodyType": text(row["body_tab"]), "review": text(row["check"])} for row in rows}


def build(source_dir: Path, body_type_csv: Path) -> dict[str, Any]:
    depth_path = source_dir / "Encar_Car_AllMakers_Depth_1004.xlsx"
    catalog_path = source_dir / "Encar_Car_Catalog_Fill_1004.xlsx"
    for path in (depth_path, catalog_path, body_type_csv):
        if not path.is_file():
            raise FileNotFoundError(path)

    depth_rows = rows_by_header(depth_path, "전체_뎁스")
    summary_rows = rows_by_header(depth_path, "제조사_요약")
    make_meta = exact_map(rows_by_header(catalog_path, "제조사"), ("제조사(엔카 원문)",), "제조사")
    group_meta = exact_map(rows_by_header(catalog_path, "모델그룹"), ("제조사", "모델그룹(엔카 원문)"), "모델그룹")
    generation_meta = exact_map(rows_by_header(catalog_path, "세부모델(세대)"), ("제조사", "모델그룹", "세부모델(엔카 원문)"), "세부모델")
    fuel_meta_rows = rows_by_header(catalog_path, "연료·구동")
    grade_meta_rows = rows_by_header(catalog_path, "등급")
    fuel_meta = exact_map(fuel_meta_rows, ("제조사", "모델그룹", "세부모델", "연료·구동(엔카 원문)"), "연료·구동")
    grade_meta = exact_map(grade_meta_rows, ("제조사", "모델그룹", "세부모델", "연료·구동", "등급(엔카 원문)"), "등급")
    body_types = load_body_types(body_type_csv)

    generation_order = orders_with_blanks_last(
        list(generation_meta.values()),
        ("제조사", "모델그룹"),
        "세부모델(엔카 원문)",
    )
    grade_order = orders_with_blanks_last(
        grade_meta_rows,
        ("제조사", "모델그룹", "세부모델", "연료·구동"),
        "등급(엔카 원문)",
    )

    fuel_order: dict[tuple[str, str, str, str], int] = {}
    fuel_parent_counts: dict[tuple[str, str, str], int] = defaultdict(int)
    for row in fuel_meta_rows:
        fuel_name = text(row["연료·구동(엔카 원문)"])
        if not fuel_name:
            continue
        parent = tuple(text(row[column]) for column in ("제조사", "모델그룹", "세부모델"))
        fuel_parent_counts[parent] += 1
        fuel_order[(*parent, fuel_name)] = fuel_parent_counts[parent]

    subgrade_order: dict[tuple[str, str, str, str, str, str], int] = {}
    subgrade_parent_counts: dict[tuple[str, str, str, str, str], int] = defaultdict(int)
    for row in depth_rows:
        subgrade_name = text(row["5뎁스 세부등급"])
        if not subgrade_name:
            continue
        parent = tuple(text(row[column]) for column in ("제조사", "1뎁스 모델그룹", "2뎁스 모델(세대)", "3뎁스 연료·구동", "4뎁스 등급"))
        full_path = (*parent, subgrade_name)
        if full_path not in subgrade_order:
            subgrade_parent_counts[parent] += 1
            subgrade_order[full_path] = subgrade_parent_counts[parent]

    depth_groups = {
        "manufacturers": group(depth_rows, ("국산/수입", "제조사")),
        "modelGroups": group(depth_rows, ("제조사", "1뎁스 모델그룹")),
        "generations": group(depth_rows, ("제조사", "1뎁스 모델그룹", "2뎁스 모델(세대)")),
        "fuelDrives": group([row for row in depth_rows if text(row["3뎁스 연료·구동"])], ("제조사", "1뎁스 모델그룹", "2뎁스 모델(세대)", "3뎁스 연료·구동")),
        "grades": group([row for row in depth_rows if text(row["4뎁스 등급"])], ("제조사", "1뎁스 모델그룹", "2뎁스 모델(세대)", "3뎁스 연료·구동", "4뎁스 등급")),
        "subgrades": group([row for row in depth_rows if text(row["5뎁스 세부등급"])], ("제조사", "1뎁스 모델그룹", "2뎁스 모델(세대)", "3뎁스 연료·구동", "4뎁스 등급", "5뎁스 세부등급")),
    }
    actual = {name: len(values) for name, values in depth_groups.items()}
    if actual != EXPECTED:
        raise ValueError(f"목표 건수 불일치: expected={EXPECTED}, actual={actual}")

    summary_by_name = {text(row["제조사"]): row for row in summary_rows}
    manufacturer_ids: dict[str, str] = {}
    model_group_ids: dict[tuple[str, str], str] = {}
    generation_ids: dict[tuple[str, str, str], str] = {}
    fuel_ids: dict[tuple[str, str, str, str], str] = {}
    grade_ids: dict[tuple[str, str, str, str, str], str] = {}

    manufacturers = []
    for (origin, name), source_rows in depth_groups["manufacturers"].items():
        meta = make_meta.get((name,))
        if not meta:
            raise ValueError(f"제조사 메타 누락: {name}")
        key = f"make_car_encar_{token(origin, name)}"
        manufacturer_ids[name] = key
        summary = summary_by_name[name]
        status, reason = review("")
        manufacturers.append({
            "key": key,
            "uuid": stable_uuid(key),
            "sourceSystem": "ENCAR",
            "sourceCode": nullable_text(meta["엔카 코드"]),
            "sourceName": source_text(meta["제조사(엔카 원문)"]),
            "displayName": name,
            "englishName": text(meta["영문명"]) or None,
            "origin": origin,
            "countryName": text(meta["국가"]) or None,
            "countryCode": text(meta["국가코드"]) or None,
            "isPopular": text(meta["인기 제조사"]).upper() in ("Y", "YES", "TRUE", "1", "인기"),
            "isVisible": name != "지리",
            "visibilityReason": "국내 정식 판매 여부 확인 필요" if name == "지리" else None,
            "sortOrder": integer(meta["엔카 정렬순서"]) or 0,
            "listingCount": unique_listing_count(source_rows, "제조사 매물", name),
            "summaryRowCount": integer(summary["행 수"]),
            "reviewStatus": status,
            "reviewReason": reason,
        })

    model_groups = []
    for path, source_rows in depth_groups["modelGroups"].items():
        make_name, name = path
        meta = group_meta.get(path)
        if not meta:
            raise ValueError(f"모델그룹 메타 누락: {' › '.join(path)}")
        key = f"model_car_encar_group_{token(make_name, name)}"
        model_group_ids[path] = key
        body = body_types.get(name) if make_name == "현대" else None
        status, reason = review(text(meta["검수"]), body["review"] if body else "")
        hidden = make_name == "지리"
        model_groups.append({
            "key": key,
            "uuid": stable_uuid(key),
            "makeKey": manufacturer_ids[make_name],
            "sourceSystem": "ENCAR",
            "sourceCode": nullable_text(meta["엔카 코드"]),
            "sourceName": source_text(meta["모델그룹(엔카 원문)"]),
            "displayName": name,
            "englishName": text(meta["영문명"]) or None,
            "level": "MODEL_GROUP",
            "bodyType": body["bodyType"] if body else None,
            "isVisible": not hidden,
            "sortOrder": integer(meta["엔카 정렬순서"]) or 0,
            "priceMin10kKrw": money(meta["시세 최저(만원)"]),
            "priceMax10kKrw": money(meta["시세 최고(만원)"]),
            "listingCount": unique_listing_count(source_rows, "모델그룹 매물", " › ".join(path)),
            "reviewStatus": status,
            "reviewReason": reason,
        })

    generations = []
    for path, source_rows in depth_groups["generations"].items():
        make_name, group_name, name = path
        meta = generation_meta.get(path)
        if not meta:
            raise ValueError(f"세부모델 메타 누락: {' › '.join(path)}")
        key = f"model_car_encar_generation_{token(make_name, group_name, name)}"
        generation_ids[path] = key
        status, reason = review(text(meta["검수"]))
        generations.append({
            "key": key,
            "uuid": stable_uuid(key),
            "makeKey": manufacturer_ids[make_name],
            "parentModelKey": model_group_ids[(make_name, group_name)],
            "sourceSystem": "ENCAR",
            "sourceCode": nullable_text(meta["엔카 코드"]),
            "sourceName": source_text(meta["세부모델(엔카 원문)"]),
            "displayName": name,
            "level": "GENERATION",
            "generationCode": text(meta["세대코드(이름에서 추출)"]) or None,
            "releaseYm": ym(meta["출시 연월"]),
            "endYm": ym(meta["단종 연월"]),
            "salesStatus": text(meta["판매상태"]) or None,
            "encarImagePath": text(meta["엔카 이미지 경로"]) or None,
            "isVisible": make_name != "지리",
            "sortOrder": generation_order[path],
            "priceMin10kKrw": money(meta["시세 최저(만원)"]),
            "priceMax10kKrw": money(meta["시세 최고(만원)"]),
            "listingCount": unique_listing_count(source_rows, "모델 매물", " › ".join(path)),
            "reviewStatus": status,
            "reviewReason": reason,
        })

    fuel_drives = []
    for path, source_rows in depth_groups["fuelDrives"].items():
        make_name, group_name, generation_name, name = path
        meta = fuel_meta.get(path)
        if not meta:
            raise ValueError(f"연료·구동 메타 누락: {' › '.join(path)}")
        key = f"trim_car_encar_fuel_{token(*path)}"
        fuel_ids[path] = key
        status, reason = review(text(meta["검수"]))
        fuel_drives.append({
            "key": key,
            "uuid": stable_uuid(key),
            "generationKey": generation_ids[(make_name, group_name, generation_name)],
            "parentTrimKey": None,
            "sourceSystem": "ENCAR",
            "sourceCode": nullable_text(meta["엔카 코드"]),
            "sourceName": source_text(meta["연료·구동(엔카 원문)"]),
            "displayName": name,
            "level": "FUEL_DRIVE",
            "valueType": text(meta["값 유형"]) or None,
            "fuel": text(meta["연료"]) or None,
            "drive": text(meta["구동"]) or None,
            "isVisible": make_name != "지리",
            "sortOrder": fuel_order[path],
            "listingCount": unique_listing_count(source_rows, "연료·구동 매물", " › ".join(path)),
            "reviewStatus": status,
            "reviewReason": reason,
        })

    grades = []
    for path, source_rows in depth_groups["grades"].items():
        make_name, group_name, generation_name, fuel_name, name = path
        meta = grade_meta.get(path)
        if not meta:
            raise ValueError(f"등급 메타 누락: {' › '.join(path)}")
        key = f"trim_car_encar_grade_{token(*path)}"
        grade_ids[path] = key
        status, reason = review(text(meta["검수"]))
        grades.append({
            "key": key,
            "uuid": stable_uuid(key),
            "generationKey": generation_ids[(make_name, group_name, generation_name)],
            "parentTrimKey": fuel_ids[(make_name, group_name, generation_name, fuel_name)],
            "sourceSystem": "ENCAR",
            "sourceCode": nullable_text(meta["엔카 코드"]),
            "sourceName": source_text(meta["등급(엔카 원문)"]),
            "displayName": name,
            "level": "GRADE",
            "isVisible": make_name != "지리",
            "sortOrder": grade_order[path],
            "priceMin10kKrw": money(meta["시세 최저(만원)"]),
            "priceMax10kKrw": money(meta["시세 최고(만원)"]),
            "listingCount": unique_listing_count(source_rows, "등급 매물", " › ".join(path)),
            "reviewStatus": status,
            "reviewReason": reason,
        })

    subgrades = []
    for path, source_rows in depth_groups["subgrades"].items():
        make_name, group_name, generation_name, fuel_name, grade_name, name = path
        key = f"trim_car_encar_subgrade_{token(*path)}"
        status, reason = review("")
        subgrades.append({
            "key": key,
            "uuid": stable_uuid(key),
            "generationKey": generation_ids[(make_name, group_name, generation_name)],
            "parentTrimKey": grade_ids[(make_name, group_name, generation_name, fuel_name, grade_name)],
            "sourceSystem": "ENCAR",
            "sourceCode": None,
            "sourceName": source_text(source_rows[0]["5뎁스 세부등급"]),
            "displayName": name,
            "level": "SUBGRADE",
            "isVisible": make_name != "지리",
            "sortOrder": subgrade_order[path],
            "listingCount": unique_listing_count(source_rows, "세부등급 매물", " › ".join(path)),
            "reviewStatus": status,
            "reviewReason": reason,
        })

    normalized = {
        "meta": {
            "schemaVersion": 1,
            "sourceSystem": "ENCAR",
            "vehicleScope": "CAR",
            "snapshotDate": SNAPSHOT_DATE,
            "listingCountKind": "ENCAR_TEST_SNAPSHOT",
            "sourceFiles": [
                {"name": depth_path.name, "sha256": sha256_file(depth_path)},
                {"name": catalog_path.name, "sha256": sha256_file(catalog_path)},
                {"name": body_type_csv.name, "sha256": sha256_file(body_type_csv)},
            ],
            "counts": EXPECTED,
            "rawCounts": {
                "depthRows": len(depth_rows),
                "catalogFuelDriveRows": len(fuel_meta_rows),
                "catalogGradeRows": len(grade_meta_rows),
            },
            "exclusions": {
                "fuelDrives": [
                    {"path": [text(row[column]) for column in ("제조사", "모델그룹", "세부모델", "연료·구동(엔카 원문)")], "reason": "연료·구동 원문명이 비어 있어 선택 노드로 만들 수 없음"}
                    for row in fuel_meta_rows if not text(row["연료·구동(엔카 원문)"])
                ],
                "grades": [
                    {"path": [text(row[column]) for column in ("제조사", "모델그룹", "세부모델", "연료·구동", "등급(엔카 원문)")], "reason": "등급 원문명이 비어 있어 선택 노드로 만들 수 없음"}
                    for row in grade_meta_rows if not text(row["등급(엔카 원문)"])
                ],
            },
            "imagePolicy": "엔카 이미지 경로는 저장만 하며 서비스 화면에서 사용하지 않음",
        },
        "manufacturers": manufacturers,
        "modelGroups": model_groups,
        "generations": generations,
        "fuelDrives": fuel_drives,
        "grades": grades,
        "subgrades": subgrades,
    }
    for name, expected in EXPECTED.items():
        if len(normalized[name]) != expected:
            raise ValueError(f"정규화 결과 건수 불일치: {name} expected={expected} actual={len(normalized[name])}")
    return normalized


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-dir", type=Path, default=Path(r"C:\Users\bobae\Downloads\Claude outputs"))
    parser.add_argument("--body-type-csv", type=Path, default=Path(r"C:\Users\bobae\Downloads\Claude outputs\grandeur_test_1005\hyundai_models_bodytype.csv"))
    parser.add_argument("--output", type=Path, default=Path(__file__).with_name("encar-car-depth.normalized.json"))
    args = parser.parse_args()
    normalized = build(args.source_dir, args.body_type_csv)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(normalized, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(args.output), "counts": normalized["meta"]["counts"], "exclusions": {key: len(value) for key, value in normalized["meta"]["exclusions"].items()}}, ensure_ascii=False))


if __name__ == "__main__":
    main()
